from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "paper_artifacts" / "results_section_completion_checklist_filled.docx"


def set_cell(cell, text, bold=False, size=8):
    cell.text = str(text)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.bold = bold
            run.font.size = Pt(size)


def add_table(doc, headers, rows, size=7.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell(cell, header, bold=True, size=size)
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            set_cell(cell, value, size=size)
    return table


def add_note(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    run.font.size = Pt(9)


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9)

    title = doc.add_heading("JusticeCompass Results Section: Completion Checklist", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph(
        "Compact reporting template for Section VII. Replace bracketed fields only with measurements from the held-out evaluation run."
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("Current Evidence Status", level=1)
    add_note(
        doc,
        "The repository currently supports implementation and corpus results, but does not contain a held-out benchmark with retrieval, expert-review, outcome-estimation, or operational metrics. Do not invent values. Report pending fields as N/R (not run) until measured.",
    )
    add_table(
        doc,
        ["Verified item", "Current value"],
        [
            ["Automated validation", "8 tests passed; 74.50 s runtime"],
            ["Knowledge base", "21 statute documents; 3,282 sections; 1,392 case-law records; 270 cross-reference entries; 5,452 vector records"],
            ["Citation coverage", "94.7% in the cross-reference layer; 90 IPC sections remain unmapped"],
            ["Safety behavior", "Returns insufficient_data when comparable outcome data are too small"],
            ["Held-out benchmark", "Not run; formal retrieval, expert-review, outcome, and load metrics are N/R"],
        ],
        size=8,
    )

    doc.add_heading("Measured Retrieval Smoke Evidence (not a benchmark)", level=1)
    add_note(doc, "Four manual queries were run with hybrid retrieval, K=5, alpha=0.6. No relevance labels were available, so ranked metrics are not computed.")
    add_table(
        doc,
        ["Query", "Collection", "Top result", "Top-3 fused scores"],
        [
            ["Punishment for murder", "statutes / IPC", "IPC 302: Punishment for murder", "0.847, 0.705, 0.702"],
            ["Bail conditions in criminal cases", "caselaw", "Metadata incomplete; no validated title", "0.600, 0.400, 0.298"],
            ["Rent control in Tamil Nadu", "statutes / TNRRRLT", "Section 47: Repeal and savings", "0.871, 0.784, 0.756"],
            ["Constitutional equality", "statutes / COI", "Article 14: Equality before law", "0.600, 0.528, 0.454"],
        ],
        size=7.5,
    )

    doc.add_heading("1. Retrieval Quality", level=1)
    add_note(doc, "Report K explicitly, use the same held-out questions for all methods, and break out domain and jurisdiction.")
    add_table(
        doc,
        ["Method", "Domain / jurisdiction", "N questions", "K", "P@K", "R@K", "Hit@K", "MRR", "nDCG@K", "95% CI"],
        [
            ["Semantic-only", "[domain / jurisdiction]", "[N]", "3", "[ ]", "[ ]", "[ ]", "[ ]", "[ ]", "[ ]"],
            ["Lexical-only", "[domain / jurisdiction]", "[N]", "3", "[ ]", "[ ]", "[ ]", "[ ]", "[ ]", "[ ]"],
            ["Hybrid", "[domain / jurisdiction]", "[N]", "3", "[ ]", "[ ]", "[ ]", "[ ]", "[ ]", "[ ]"],
            ["Semantic-only", "[repeat for K=5 and 10]", "[N]", "5 / 10", "[ ]", "[ ]", "[ ]", "[ ]", "[ ]", "[ ]"],
            ["Lexical-only", "[repeat for K=5 and 10]", "[N]", "5 / 10", "[ ]", "[ ]", "[ ]", "[ ]", "[ ]", "[ ]"],
            ["Hybrid", "[repeat for K=5 and 10]", "[N]", "5 / 10", "[ ]", "[ ]", "[ ]", "[ ]", "[ ]", "[ ]"],
        ],
    )

    doc.add_heading("2. Answer Quality and Groundedness", level=1)
    add_note(doc, "State the number of reviewed responses, reviewer count, scoring scale, agreement statistic, and hallucinated-citation rate.")
    add_table(
        doc,
        ["Domain / case type", "N responses", "Reviewers", "Cit. P", "Cit. R", "Complete", "Grounded", "Relevant", "Actionable", "Safety", "Valid schema", "Fab. cit.", "Kappa"],
        [
            ["[domain / case type]", "[N]", "[N]", "[%]", "[%]", "[%]", "[%]", "1-5", "1-5", "[%]", "[%]", "[%]", "[ ]"],
            ["All / weighted", "[N]", "[N]", "[%]", "[%]", "[%]", "[%]", "1-5", "1-5", "[%]", "[%]", "[%]", "[ ]"],
        ],
        size=6.8,
    )
    add_note(doc, "Required note: citation precision = supported citations / cited citations; citation recall = required authoritative sources cited / required sources; structured validity = parses without manual repair / all responses.")

    doc.add_heading("3. Historical Outcome Estimation", level=1)
    add_note(doc, "Give train/test counts and class distribution before reporting scores. Include a majority-class baseline and a reliability diagram in the paper or supplement.")
    add_table(
        doc,
        ["Category / jurisdiction", "Train N", "Test N", "Allowed", "Dismissed", "Settled", "Other", "Majority baseline"],
        [
            ["[category / jurisdiction]", "[N]", "[N]", "[n / %]", "[n / %]", "[n / %]", "[n / %]", "[accuracy]"],
            ["[next category]", "[N]", "[N]", "[n / %]", "[n / %]", "[n / %]", "[n / %]", "[accuracy]"],
        ],
        size=7,
    )
    add_table(
        doc,
        ["Category / jurisdiction", "Accuracy", "Per-class P/R/F1", "Macro-F1", "Micro-F1", "ROC-AUC", "PR-AUC", "Brier", "ECE", "Coverage", "Abstention", "95% CI"],
        [
            ["[category / jurisdiction]", "[ ]", "[allowed: / / ; ...]", "[ ]", "[ ]", "[ ]", "[ ]", "[ ]", "[ ]", "[%]", "[%]", "[ ]"],
            ["[next category]", "[ ]", "[allowed: / / ; ...]", "[ ]", "[ ]", "[ ]", "[ ]", "[ ]", "[ ]", "[%]", "[%]", "[ ]"],
        ],
        size=6.8,
    )
    add_note(doc, "Figure required: reliability diagram with predicted-probability bins, observed outcome frequency, sample count per bin, and the ideal diagonal. Report abstention as insufficient_data responses / all eligible questions.")

    doc.add_heading("4. Operational Reliability", level=1)
    add_note(doc, "Record the workload, measurement period, hardware/infra, concurrency, model, and cache configuration alongside every value.")
    add_table(
        doc,
        ["Condition", "Queries", "Concurrency", "p50", "p95", "p99", "Error rate", "Timeout rate", "Cache hit", "API calls / req", "Cost / req", "Docs / req", "Tokens / req"],
        [
            ["[hardware, model, date range]", "[N]", "[N]", "[ms]", "[ms]", "[ms]", "[%]", "[%]", "[%]", "[N]", "[$]", "[N]", "[N]"],
        ],
        size=6.8,
    )

    doc.add_heading("Dataset and Concrete Example", level=1)
    add_table(
        doc,
        ["Metadata to report", "Value"],
        [
            ["Total evaluation questions / outcome records", "[N questions] / [N records]"],
            ["Domains and jurisdictions", "[list; include per-group counts]"],
            ["Split and date range", "[train / validation / held-out; temporal cutoff]"],
            ["Model, prompt, retrieval K, threshold", "[values]"],
            ["Confidence interval method", "[bootstrap / Wilson / other]"],
        ],
        size=8,
    )
    doc.add_paragraph("Concrete response example to include in Section VII:")
    add_table(
        doc,
        ["Question", "Answer excerpt", "Citations", "Retrieval scores", "Outcome score / abstention", "Reviewer notes"],
        [["[real held-out question]", "[short answer]", "[source IDs / sections]", "[top-k scores]", "[value or insufficient_data]", "[grounding / safety]"],
         ["", "", "", "", "", ""]],
        size=7,
    )

    doc.add_heading("Paper Decision Rule", level=1)
    add_note(doc, "Section VII is complete only when every table has numeric values, denominators, subgroup labels, uncertainty estimates, and test conditions. Until then, describe the work as implementation validation and corpus coverage, not as a completed benchmark evaluation.")

    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if "N/R" in run.text or "pending" in run.text.lower():
                run.font.color.rgb = RGBColor(156, 64, 45)

    OUTPUT.parent.mkdir(exist_ok=True)
    doc.save(OUTPUT)
    print(f"Saved {OUTPUT}")


if __name__ == "__main__":
    build_document()